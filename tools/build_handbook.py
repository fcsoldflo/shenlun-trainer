# -*- coding: utf-8 -*-
"""从 book.json 生成《申论全书 · 速记手册》DOCX。"""
import json
import os
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INK = RGBColor(0x21, 0x1D, 0x18)
RED = RGBColor(0x9C, 0x2B, 0x23)
GOLD = RGBColor(0x8A, 0x6A, 0x2F)
GREY = RGBColor(0x6F, 0x64, 0x55)
LIGHT = RGBColor(0xA3, 0x96, 0x7E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "微软雅黑"
SERIF = "宋体"

BASE = os.path.dirname(os.path.abspath(__file__))
with open(os.path.join(BASE, "book.json"), "r", encoding="utf-8") as f:
    BOOK = json.load(f)
CHAPTERS = BOOK["chapters"]

doc = Document()

# ---------- page ----------
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(1.35)
sec.bottom_margin = Cm(1.25)
sec.left_margin = Cm(1.5)
sec.right_margin = Cm(1.5)
sec.header_distance = Cm(0.7)
sec.footer_distance = Cm(0.6)
CONTENT_W = 18.0

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(8.8)
normal.font.color.rgb = INK
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
pf = normal.paragraph_format
pf.space_before = Pt(0)
pf.space_after = Pt(2)
pf.line_spacing = 1.1


def set_font(run, name=FONT, size=8.8, bold=None, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def run_shade(run, color):
    rPr = run._element.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    rPr.append(shd)


def cell_shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def para_shade(p, color):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    pPr.append(shd)


def para_border(p, edge="left", color="9C2B23", sz=12, space=4):
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    pBdr.append(el)


def table_borders(table, color="DCD2C0", sz=4):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def cell_margins(table, top=40, bottom=40, left=75, right=75):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for tag, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def set_widths(table, widths):
    table.autofit = False
    for i, w in enumerate(widths):
        table.columns[i].width = Cm(w)
    for row in table.rows:
        for i, w in enumerate(widths):
            if i < len(row.cells):
                row.cells[i].width = Cm(w)
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(int(sum(widths) / 2.54 * 1440)))


def row_no_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trPr.append(el)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    trPr.append(el)


def cell_write(cell, text, size=8.2, bold=False, color=INK, name=FONT,
               align=WD_ALIGN_PARAGRAPH.LEFT, line=1.08):
    cell.text = ""
    lines = str(text).split("\n")
    for i, line_text in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = line
        r = p.add_run(line_text)
        set_font(r, name=name, size=size, bold=bold, color=color)


def add_para(text="", size=8.8, bold=False, color=INK, name=FONT,
             before=0, after=2, line=1.1, align=None, indent=0, hanging=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if align is not None:
        p.alignment = align
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    if hanging:
        p.paragraph_format.first_line_indent = Cm(-hanging)
    if text:
        r = p.add_run(text)
        set_font(r, name=name, size=size, bold=bold, color=color)
    return p


def section_heading(num, title, sub=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    r0 = p.add_run(" " + str(num) + " ")
    set_font(r0, name=SERIF, size=9, bold=True, color=WHITE)
    run_shade(r0, "9C2B23")
    r1 = p.add_run("  " + title)
    set_font(r1, size=11.5, bold=True, color=INK)
    if sub:
        r2 = p.add_run("   " + sub)
        set_font(r2, size=7.8, color=LIGHT)
    return p


def render_kv(block):
    rows = block.get("rows", [])
    t = doc.add_table(rows=len(rows), cols=2)
    table_borders(t)
    cell_margins(t)
    set_widths(t, [3.0, 15.0])
    for i, (k, v) in enumerate(rows):
        cell_shade(t.cell(i, 0), "FAF5E9")
        cell_write(t.cell(i, 0), k, size=8.4, bold=True, color=RGBColor(0x5C, 0x50, 0x40))
        cell_write(t.cell(i, 1), v, size=8.5)
        row_no_split(t.rows[i])
    add_para("", size=2, after=3)


def render_table(block):
    head = block.get("head", [])
    rows = block.get("rows", [])
    ncol = len(head) if head else max(len(r) for r in rows)
    if ncol == 2:
        widths = [3.2, 14.8]
    elif ncol == 3:
        widths = [3.6, 6.6, 7.8]
    elif ncol == 4:
        widths = [3.2, 4.8, 4.8, 5.2]
    elif ncol == 6:
        widths = [2.4, 5.8, 2.4, 2.4, 2.4, 2.4]
    else:
        widths = [CONTENT_W / ncol] * ncol

    t = doc.add_table(rows=1 + len(rows), cols=ncol)
    table_borders(t)
    cell_margins(t, top=35, bottom=35, left=60, right=60)
    set_widths(t, widths)
    if head:
        for j, h in enumerate(head):
            c = t.cell(0, j)
            cell_shade(c, "211D18")
            cell_write(c, h, size=7.9, bold=True, color=WHITE)
        repeat_header(t.rows[0])
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            if j == 0:
                cell_shade(c, "FAF5E9")
                cell_write(c, val, size=7.9, bold=True, color=RGBColor(0x5C, 0x50, 0x40))
            else:
                cell_write(c, val, size=7.9)
        row_no_split(t.rows[i])
    add_para("", size=2, after=3)


def render_note(block):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Cm(0.15)
    p.paragraph_format.right_indent = Cm(0.1)
    para_shade(p, "FFFBF0")
    para_border(p, "left", "9C2B23", 12, 5)
    if block.get("title"):
        r = p.add_run(block["title"] + "：")
        set_font(r, size=8.6, bold=True, color=RED)
    r2 = p.add_run(block.get("body", "").replace("\n", "  "))
    set_font(r2, size=8.7)


def render_list(block):
    for i, item in enumerate(block.get("items", [])):
        p = add_para("• " + item, size=8.6, after=1, indent=0.45, hanging=0.35)
    if block.get("after"):
        add_para(block["after"], size=7.8, color=LIGHT, after=3)


def render_grid(block):
    items = block.get("items", [])
    ncol = block.get("cols", 2 if len(items) % 2 == 0 else 1)
    if len(items) <= 1:
        ncol = 1
    rows = (len(items) + ncol - 1) // ncol
    t = doc.add_table(rows=rows, cols=ncol)
    table_borders(t, color="E6DCCB")
    cell_margins(t, top=45, bottom=45, left=80, right=80)
    set_widths(t, [CONTENT_W / ncol] * ncol)
    for idx, it in enumerate(items):
        r, c = idx // ncol, idx % ncol
        cell = t.cell(r, c)
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        rr = p.add_run(it.get("t", ""))
        set_font(rr, size=8.5, bold=True, color=RED)
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.line_spacing = 1.08
        r2 = p2.add_run(it.get("b", ""))
        set_font(r2, size=8.2)
    add_para("", size=2, after=3)


def render_elements(block):
    for it in block.get("items", []):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(it.get("name", ""))
        set_font(r, name=SERIF, size=10.5, bold=True, color=RED)
        for line in it.get("lines", []):
            q = doc.add_paragraph()
            q.paragraph_format.space_before = Pt(0)
            q.paragraph_format.space_after = Pt(0.5)
            q.paragraph_format.line_spacing = 1.12
            q.paragraph_format.left_indent = Cm(0.2)
            if "：" in line:
                k, v = line.split("：", 1)
                r1 = q.add_run(k + "：")
                set_font(r1, size=8.5, bold=True, color=RGBColor(0x5C, 0x50, 0x40))
                r2 = q.add_run(v)
                set_font(r2, size=8.5)
            else:
                r3 = q.add_run(line)
                set_font(r3, size=8.5)


def render_quote(block):
    t = doc.add_table(rows=1, cols=1)
    table_borders(t, color="E8D6B8")
    cell_margins(t, top=60, bottom=60, left=100, right=100)
    set_widths(t, [CONTENT_W])
    c = t.cell(0, 0)
    cell_shade(c, "FFF8EA")
    cell_write(c, block.get("body", ""), size=9.4, name=SERIF,
               align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para("", size=2, after=3)


def render_block(block):
    t = block.get("type")
    if t == "kv":
        render_kv(block)
    elif t == "table":
        render_table(block)
    elif t == "note":
        render_note(block)
    elif t == "list":
        render_list(block)
    elif t == "grid":
        render_grid(block)
    elif t == "elements":
        render_elements(block)
    elif t == "quote":
        render_quote(block)
    elif t == "chips":
        text = "　".join(block.get("items", []))
        note = {"type": "note", "title": block.get("title", ""), "body": text}
        render_note(note)


def add_footer():
    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("申论全书 · 速记手册　|　第 ")
    set_font(r, size=7.2, color=LIGHT)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    rr = OxmlElement("w:r")
    tt = OxmlElement("w:t")
    tt.text = "1"
    rr.append(tt)
    fld.append(rr)
    p._p.append(fld)
    r2 = p.add_run(" 页")
    set_font(r2, size=7.2, color=LIGHT)


# ================= 封面 =================
bar = doc.add_table(rows=1, cols=1)
table_borders(bar, color="9C2B23", sz=4)
set_widths(bar, [CONTENT_W])
bc = bar.cell(0, 0)
cell_shade(bc, "9C2B23")
bc.text = ""
bp = bc.paragraphs[0]
bp.paragraph_format.space_before = Pt(10)
bp.paragraph_format.space_after = Pt(10)
br = bp.add_run("全书七章 · 含补充题型")
set_font(br, size=10, bold=True, color=WHITE)

add_para("", size=20, after=30)
add_para("申论全书", size=34, bold=True, color=INK, name=SERIF, after=2)
add_para("速记手册", size=24, bold=True, color=RED, name=SERIF, after=14)
add_para("概括归纳 · 对策启示 · 综合分析 · 公文写作 · 作文 · 补充题型 · 材料阅读",
         size=9.5, color=GREY, after=26)

stat = doc.add_table(rows=1, cols=4)
table_borders(stat, color="E6DCCB")
cell_margins(stat, top=70, bottom=70)
set_widths(stat, [4.5, 4.5, 4.5, 4.5])
stats = [
    (str(len(CHAPTERS)), "章"),
    (str(sum(len(c["cards"]) for c in CHAPTERS)), "张速记卡"),
    (str(sum(len(c["cloze"]) for c in CHAPTERS)), "句填空"),
    (str(sum(len(c["quiz"]) for c in CHAPTERS)), "道测验"),
]
for i, (n, label) in enumerate(stats):
    c = stat.cell(0, i)
    c.text = ""
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(n)
    set_font(r, name=SERIF, size=18, bold=True, color=RED)
    p2 = c.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(label)
    set_font(r2, size=7.5, color=GREY)

add_para("", size=20, after=40)
add_para("目录", size=11, bold=True, color=RED, name=SERIF, after=4)
for ch in CHAPTERS:
    p = add_para("", after=1)
    r1 = p.add_run("第 " + ch["no"] + " 章　")
    set_font(r1, size=9, color=GOLD, bold=True)
    r2 = p.add_run(ch["title"])
    set_font(r2, size=9.5, bold=True)
    r3 = p.add_run("　　" + ch["sub"])
    set_font(r3, size=7.8, color=LIGHT)

add_para("", size=16, after=20)
add_para("依据讲义原文整理 ｜ 口诀、测验与排版为辅助学习自编 ｜ 源文件缺口见末页",
         size=7.5, color=LIGHT)

doc.add_page_break()

# ================= 各章 =================
for ci, ch in enumerate(CHAPTERS):
    if ci > 0:
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_before = Pt(4)
        spacer.paragraph_format.space_after = Pt(10)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    r = p.add_run("第 " + ch["no"] + " 章")
    set_font(r, size=8.5, bold=True, color=GOLD)
    t = add_para(ch["title"], size=19, bold=True, name=SERIF, after=1)
    t.paragraph_format.keep_with_next = True
    s = add_para(ch["sub"], size=8.6, color=GREY, after=2)
    s.paragraph_format.keep_with_next = True
    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(0)
    rule.paragraph_format.space_after = Pt(3)
    rule.paragraph_format.keep_with_next = True
    para_border(rule, "bottom", "211D18", 8, 2)

    for si, section in enumerate(ch.get("sections", []), start=1):
        sh = section_heading(si, section.get("h", ""), section.get("sub", ""))
        sh.paragraph_format.keep_with_next = True
        for block in section.get("blocks", []):
            render_block(block)

    if ch.get("gaps"):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Cm(0.15)
        para_shade(p, "FBF5E7")
        r = p.add_run("本章原文缺口：")
        set_font(r, size=7.8, bold=True, color=RGBColor(0x8A, 0x5A, 0x1F))
        r2 = p.add_run("；".join(g.rstrip("。") for g in ch["gaps"]) + "。")
        set_font(r2, size=7.6, color=GREY)

# ================= 说明页 =================
add_para("", size=6, after=6)
add_para("说明与勘误", size=16, bold=True, name=SERIF, after=2)
add_para("这份手册哪里可靠、哪里存疑，一次讲清楚。", size=8.5, color=GREY, after=8)

errata = [
    ("源文件情况", [
        "第一份《第一章 概括归纳很精.docx》：24 页，实际含五章＋补充题型，文件名与内容不符；0 张图片，"
        "正文有图片转写痕迹（工具生成语句、$$...$$ 占位符）。",
        "第二份《启承申论笔记.docx》：109 页、124 张图，按正文顺序抽出 116 张（110 张整页＋6 张裁剪）；"
        "文字层仅约 6,759 字，主体是拍照/截图，含大量手写批注。",
        "本手册新增的“材料阅读”章，以及对策问法矩阵、大公文、分论点三法等，来自第二份文件的逐页读图；"
        "缺页/模糊页只取独立知识点，文章不补。",
        "缺失的例题与答案没有补写，也没有替讲师编评分点。"
    ]),
    ("具体存疑点", [
        "第二章“身份定位（×）”“可行性（×）”后的 × 含义不明，未写成结论。",
        "第四章“（注：≤60 字时，第二段占比最大）”位置在“称谓”之后，指代不明，未纳入卡片。",
        "第五章“背景 / 咔咔咔”中的“咔咔咔”疑似课程内部代称，原文如此，未擅自解释。",
        "第五章“分论点三法”已从后半本补全；“两个分论点如何补第三个”也已录入。",
        "第四章“逻辑/层次/观点/分类”落点表（P043）手写较密、辨读存疑，暂未录入。",
        "第一章“保留核心键（键词多为名词）”应为“核心词/关键词”，已按笔误处理。",
        "第三章“句子类”开头残留工具生成语句，已删除。"
    ]),
    ("编写口径", [
        "速记卡、填空、测验只使用原文明确出现的表述；推断性内容会标注或不写。",
        "口诀、测验解析、排版为自编助记，不等于讲义原文。",
        "标点、序号、明显笔误做了统一；没有改动任何结论性内容。"
    ])
]
for title, items in errata:
    h = add_para(title, size=10.5, bold=True, color=RED, name=SERIF, before=8, after=2)
    for it in items:
        add_para("• " + it, size=8.2, after=1, indent=0.4, hanging=0.3)

add_para("", size=6, after=6)
add_para("使用建议：先读各章“速览”建立框架；再用网页版速记卡自评、填空默写、闯关测验；"
         "考前只看本手册的黑体标题与红色关键词复述。", size=8, color=GREY)

add_footer()

out = os.path.join(BASE, "..", "申论全书_速记手册.docx")
out = os.path.abspath(out)
doc.save(out)
print("saved:", out)
